from io import BytesIO

from datetime import timedelta

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from accounts.roles import ROLE_ADMIN, ROLE_EDITOR, ROLE_VIEWER
from assets.models import Asset
from audit.context import audit_context
from audit.models import AuditLog
from maintenance.models import MaintenancePlan
from qualification.models import QualificationPlan
from tasks.models import Task


class AuditTrailTests(TestCase):
    def setUp(self):
        user_model = get_user_model()
        self.admin_group, _ = Group.objects.get_or_create(name=ROLE_ADMIN)
        self.editor_group, _ = Group.objects.get_or_create(name=ROLE_EDITOR)
        self.viewer_group, _ = Group.objects.get_or_create(name=ROLE_VIEWER)

        self.admin_user = user_model.objects.create_user(username="admin", password="pass-12345")
        self.editor_user = user_model.objects.create_user(username="editor", password="pass-12345")
        self.viewer_user = user_model.objects.create_user(username="viewer", password="pass-12345")

        self.admin_user.groups.add(self.admin_group)
        self.editor_user.groups.add(self.editor_group)
        self.viewer_user.groups.add(self.viewer_group)
        self.admin_user.profile.user_code = "AD"
        self.admin_user.profile.role = ROLE_ADMIN
        self.admin_user.profile.save()
        self.editor_user.profile.user_code = "ED"
        self.editor_user.profile.role = ROLE_EDITOR
        self.editor_user.profile.save()
        self.viewer_user.profile.user_code = "VW"
        self.viewer_user.profile.role = ROLE_VIEWER
        self.viewer_user.profile.save()

        self.asset = Asset.objects.create(
            asset_id="A-100",
            name="Mischer",
            commissioning_date=timezone.localdate() - timedelta(days=30),
            status=Asset.STATUS_ACTIVE,
        )
        self.maintenance_plan = MaintenancePlan.objects.create(
            asset=self.asset,
            title="Wartung Baseline",
            interval_value=30,
            interval_unit=MaintenancePlan.INTERVAL_DAYS,
            warning_days=5,
            is_active=True,
        )
        self.qualification_plan = QualificationPlan.objects.create(
            asset=self.asset,
            title="Qualifizierung Baseline",
            interval_value=60,
            interval_unit=QualificationPlan.INTERVAL_DAYS,
            warning_days=7,
            is_active=True,
        )
        self.task = Task.objects.create(
            title="Sichtprüfung durchführen",
            asset=self.asset,
            due_date=timezone.localdate() + timedelta(days=2),
            priority=Task.PRIORITY_MEDIUM,
            status=Task.STATUS_OPEN,
            responsible_user=self.editor_user,
        )

        AuditLog.objects.all().delete()

    def test_creation_logs_are_recorded(self):
        with audit_context(user=self.admin_user, change_reason="Neue Anlage übernommen"):
            asset = Asset.objects.create(asset_id="A-200", name="Abfülllinie", status=Asset.STATUS_ACTIVE)

        entry = AuditLog.objects.get(model_name="Asset", object_id=str(asset.pk), action=AuditLog.ACTION_CREATE)
        self.assertEqual(entry.user, self.admin_user)
        self.assertEqual(entry.change_reason, "Neue Anlage übernommen")
        self.assertIn("Asset-ID: A-200", entry.new_value)
        self.assertIn("Bezeichnung: Abfülllinie", entry.new_value)
        self.assertIn("Asset-ID: A-200", entry.new_value_display)
        self.assertIn("Bezeichnung: Abfülllinie", entry.new_value_display)
        self.assertEqual(entry.user_display, "admin [AD] - Admin")

    def test_update_logs_detect_changed_fields_and_status_changes(self):
        with audit_context(user=self.editor_user, change_reason="Status angepasst"):
            self.task.title = "Sichtprüfung gestartet"
            self.task.status = Task.STATUS_IN_PROGRESS
            self.task.save()

        status_entry = AuditLog.objects.get(
            model_name="Task",
            object_id=str(self.task.pk),
            field_name="status",
        )
        title_entry = AuditLog.objects.get(
            model_name="Task",
            object_id=str(self.task.pk),
            field_name="title",
        )

        self.assertEqual(status_entry.action, AuditLog.ACTION_STATUS_CHANGE)
        self.assertEqual(status_entry.old_value, "Offen")
        self.assertEqual(status_entry.new_value, "In Bearbeitung")
        self.assertEqual(status_entry.user, self.editor_user)
        self.assertEqual(title_entry.action, AuditLog.ACTION_UPDATE)
        self.assertEqual(title_entry.old_value, "Sichtprüfung durchführen")
        self.assertEqual(title_entry.new_value, "Sichtprüfung gestartet")
        self.assertEqual(status_entry.field_label, "Status")
        self.assertEqual(title_entry.field_label, "Titel")
        self.assertEqual(status_entry.user_display, "editor [ED] - Benutzer")

    def test_existing_summary_values_are_rendered_with_readable_labels(self):
        entry = AuditLog.objects.create(
            user=self.admin_user,
            action=AuditLog.ACTION_CREATE,
            model_name="SystemSettings",
            object_id="1",
            object_repr="Systemeinstellungen",
            new_value=(
                "singleton_enforcer: Ja; "
                "default_maintenance_warning_days: 7; "
                "default_maintenance_interval_value: 30; "
                "default_maintenance_interval_unit: Tage; "
                "default_qualification_warning_days: 14"
            ),
        )

        self.assertNotIn("singleton_enforcer", entry.new_value_display)
        self.assertIn("Standard Warnungstage Wartung: 7", entry.new_value_display)
        self.assertIn("Standard Intervallwert Wartung: 30", entry.new_value_display)
        self.assertIn("Standard Intervall-Einheit Wartung: Tage", entry.new_value_display)
        self.assertIn("Standard Warnungstage Qualifizierung: 14", entry.new_value_display)

    def test_request_updates_use_authenticated_user_for_audit(self):
        self.client.force_login(self.editor_user)
        response = self.client.post(
            reverse("tasks:edit", args=[self.task.pk]),
            {
                "title": self.task.title,
                "description": "",
                "asset": self.asset.pk,
                "due_date": self.task.due_date.strftime("%Y-%m-%d"),
                "priority": self.task.priority,
                "status": Task.STATUS_DONE,
                "responsible_user": self.editor_user.pk,
            },
        )

        self.assertRedirects(response, reverse("tasks:detail", args=[self.task.pk]))
        entry = AuditLog.objects.filter(model_name="Task", object_id=str(self.task.pk), field_name="status").latest("id")
        self.assertEqual(entry.user, self.editor_user)
        self.assertEqual(entry.action, AuditLog.ACTION_STATUS_CHANGE)

    def test_audit_list_requires_login(self):
        response = self.client.get(reverse("audit:list"))
        expected = f"{reverse('accounts:login')}?next={reverse('audit:list')}"
        self.assertRedirects(response, expected)

    def test_all_roles_can_access_audit_list(self):
        with audit_context(user=self.admin_user):
            self.asset.status = Asset.STATUS_INACTIVE
            self.asset.save()

        for user in (self.admin_user, self.editor_user, self.viewer_user):
            with self.subTest(user=user.username):
                self.client.force_login(user)
                response = self.client.get(reverse("audit:list"))
                self.assertEqual(response.status_code, 200)
                self.assertContains(response, "Audit Trail")

    def test_audit_list_filters_by_model_and_action(self):
        with audit_context(user=self.admin_user):
            self.asset.status = Asset.STATUS_OUT_OF_SERVICE
            self.asset.save()
            self.task.title = "Andere Maßnahme"
            self.task.save()

        self.client.force_login(self.viewer_user)
        response = self.client.get(
            reverse("audit:list"),
            {"model": "Asset", "action": AuditLog.ACTION_STATUS_CHANGE},
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.asset.asset_id)
        self.assertNotContains(response, self.task.title)

    def test_object_histories_appear_on_detail_pages(self):
        with audit_context(user=self.admin_user):
            self.asset.status = Asset.STATUS_INACTIVE
            self.asset.save()
            self.maintenance_plan.responsible_person = "Team Blau"
            self.maintenance_plan.save()
            self.qualification_plan.responsible_person = "QA Team"
            self.qualification_plan.save()
            self.task.status = Task.STATUS_IN_PROGRESS
            self.task.save()

        self.client.force_login(self.viewer_user)

        asset_response = self.client.get(reverse("assets:detail", args=[self.asset.pk]))
        maintenance_response = self.client.get(
            reverse("maintenance:plan-detail", args=[self.maintenance_plan.pk])
        )
        qualification_response = self.client.get(
            reverse("qualification:plan-detail", args=[self.qualification_plan.pk])
        )
        task_response = self.client.get(reverse("tasks:detail", args=[self.task.pk]))

        self.assertContains(asset_response, "Änderungsverlauf")
        self.assertContains(asset_response, "Aktiv")
        self.assertContains(asset_response, "Inaktiv")
        self.assertContains(maintenance_response, "Team Blau")
        self.assertContains(qualification_response, "QA Team")
        self.assertContains(task_response, "In Bearbeitung")

    def test_viewer_can_export_filtered_audit_log_as_xlsx(self):
        with audit_context(user=self.admin_user):
            self.asset.status = Asset.STATUS_OUT_OF_SERVICE
            self.asset.save()
            self.task.title = "Andere Maßnahme"
            self.task.save()

        self.client.force_login(self.viewer_user)
        response = self.client.get(
            reverse("audit:list"),
            {"model": "Asset", "action": AuditLog.ACTION_STATUS_CHANGE, "export": "xlsx"},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active

        self.assertEqual(sheet["A1"].value, "Zeitpunkt")
        self.assertEqual(sheet["C2"].value, "Statuswechsel")
        self.assertEqual(sheet["D2"].value, "Asset")
        self.assertNotEqual(sheet.max_row, 1)
